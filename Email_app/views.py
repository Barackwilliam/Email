from django.shortcuts import render, redirect
from django.core.paginator import Paginator
from django.template.loader import render_to_string
from django.utils.html import strip_tags
from django.core.mail import send_mail
from django.conf import settings
from django.contrib.auth.forms import UserCreationForm, AuthenticationForm
from django.contrib.auth import login, authenticate, logout
from django.contrib.auth.decorators import login_required
from .models import EmailEntry
from .forms import EmailEntryForm
from django.db.models import Q
from django.shortcuts import render, redirect
from django.http import HttpResponse
import csv
from openpyxl import Workbook
from xhtml2pdf import pisa
import io
from .models import SMSMessage
from .forms import SMSMessageForm



from datetime import datetime


@login_required
def dashboard(request):
    current_hour = datetime.now().hour
    if 5 <= current_hour < 12:
        greeting = "Good Morning"
    elif 12 <= current_hour < 17:
        greeting = "Good Afternoon"
    elif 17 <= current_hour < 22:
        greeting = "Good Evening"
    else:
        greeting = "Good Night"

    search_query = request.GET.get('search', '')
    per_page = int(request.GET.get('per_page', 6))  # default ni 6

    if search_query:
        emails = EmailEntry.objects.filter(
            Q(email__icontains=search_query) |
            Q(region__icontains=search_query) |
            Q(district__icontains=search_query) |
            Q(school_name__icontains=search_query) |
            Q(phone_number__icontains=search_query)
        )
    else:
        emails = EmailEntry.objects.all()

    total_emails = EmailEntry.objects.count()
    filtered_count = emails.count()

    paginator = Paginator(emails, per_page)
    page_number = request.GET.get('page')
    page_obj = paginator.get_page(page_number)

    return render(request, 'dashboard.html', {
        'page_obj': page_obj,
        'total_emails': total_emails,
        'filtered_count': filtered_count,
        'greeting': greeting,
        'search_query': search_query,  # <-- Hii ndiyo ilikuwa haipo!
    })



from django.contrib import messages
import random

success_messages = [
    "Great work, {username}! Your data has been successfully saved. Your contribution to the team is highly appreciated.",
    "Thank you, {username}. The Message has been added successfully. Your collaboration is the cornerstone of our success.",
    "Well done, {username}! You have successfully added the data. We commend your consistent efforts.",
    "The data has been successfully saved. Thank you, {username}, for being part of our collective progress.",
    "Outstanding job, {username}! Keep up the great work – the team depends on you.",
    "Thank you for your contribution, {username}. You have successfully added a new Message . Together, we are building a robust system.",
    "Your submission was successful. Thank you, {username}, for your wholehearted commitment to this task.",
    "Amazing work, {username}. The data has been saved successfully. We appreciate your hard work and dedication.",
    "Excellent job, {username}. The Message  has been added to the system. We’re grateful for your contribution.",
    "Your contribution is invaluable, {username}. The data has been saved, and we appreciate your commitment to excellence.",
    "Thank you for your dedication, {username}. The data has been added successfully. The team values your efforts.",
    "Great work, {username}! You’ve successfully submitted the data. We’re grateful for your efforts in making this a success.",
    "Well done, {username}. The Message  has been successfully added to the system. Your hard work is truly appreciated.",
    "Excellent contribution, {username}. Your data has been successfully recorded. We’re proud of your dedication.",
    "Thank you for your effort, {username}. The data has been added successfully. Keep up the excellent work.",
    "Great job, {username}. Your data has been added to the system. We are thankful for your contributions.",
    "You’ve done an excellent job, {username}. Your data has been successfully added to the system. We appreciate your hard work.",
    "Kudos, {username}. Your data has been successfully submitted. Thank you for your hard work and dedication.",
    "Well done, {username}. The data has been saved successfully. Keep up the great work – the team is counting on you.",
    "Fantastic job, {username}! You’ve successfully contributed. The team appreciates your hard work and effort.",
    "Great job, {username}. The data has been successfully saved. Your contribution is valued by the entire team.",
    "Thank you for your effort, {username}. The Message  has been successfully added. Your work makes a difference.",
    "Awesome work, {username}! Your contribution has been successfully saved. The team truly appreciates your dedication.",
    "Congratulations, {username}. You’ve added data successfully. We value your contribution to the project.",
    "Nice work, {username}. Your data has been successfully saved. The team is thankful for your hard work.",
    "Great work, {username}. The data has been successfully added. Your commitment to excellence is inspiring.",
    "Thank you, {username}. The Message  has been added successfully. We’re proud of your continuous effort.",
    "Amazing effort, {username}. The data has been successfully saved. Your work has made a significant impact.",
    "You did a fantastic job, {username}. Your data has been saved. Keep up the excellent work!",
    "Fantastic job, {username}. The Message  has been successfully added. The team is grateful for your dedication.",
    "Well done, {username}. The data has been successfully saved. Your contribution is helping us move forward.",
    "Great work, {username}. Your Message  has been successfully saved. Keep up the great work, and thank you for your effort.",
    "Excellent job, {username}. The data has been saved. Your contribution is essential to our progress.",
    "Congratulations, {username}. Your Message  has been successfully recorded. The team is grateful for your hard work.",
    "Well done, {username}. The data has been added successfully. We’re excited about the progress we’re making together.",
    "Fantastic work, {username}. The Message ta has been successfully saved. The team is proud of what you’ve achieved.",
    "Thank you, {username}. Your contribution has been recorded successfully. Keep up the amazing work.",
    "Great work, {username}. You’ve successfully added the Message . The team values your continued effort.",
    "You did an excellent job, {username}. Your data has been saved successfully. We’re grateful for your contribution.",
    "Awesome work, {username}. The Message  has been added successfully. Your efforts are helping build a better system.",
    "Nice work, {username}. Your Message  has been successfully added. The team appreciates your effort and dedication.",
    "Congratulations, {username}. You’ve successfully added the data. Your dedication is helping us grow.",
    "Outstanding job, {username}. The Message  has been saved successfully. We’re thankful for your consistent contributions.",
    "Well done, {username}. Your Message  has been successfully added. Keep pushing forward – the team is behind you.",
    "Fantastic job, {username}. The data has been saved successfully. You’re making a significant difference with your work.",
    "Excellent work, {username}. Your Message  has been successfully saved. We appreciate your consistent dedication.",
    "Great job, {username}. The data has been successfully added. Thank you for your continued efforts.",
    "Awesome effort, {username}. The Message  has been added successfully. Your hard work is appreciated.",
    "Well done, {username}. The data has been successfully saved. We’re thankful for your contribution to the system.",
    "Congratulations, {username}. Your Message  has been added successfully. Your efforts are driving the project forward.",
    "Thank you for your contribution, {username}. The Message  has been successfully recorded. Keep up the great work."
]



@login_required
def add_email(request):
    if request.method == 'POST':
        form = EmailEntryForm(request.POST)
        if form.is_valid():
            email_entry = form.save(commit=False)
            email_entry.user = request.user
            email_entry.save()

            # Chagua ujumbe mmoja randomly
            msg = random.choice(success_messages).format(username=request.user.username)
            messages.success(request, msg)

            form = EmailEntryForm()  # Weka form mpya tupu
    else:
        form = EmailEntryForm()

    return render(request, 'add_email.html', {'form': form})

def register(request):
    if request.method == 'POST':
        form = UserCreationForm(request.POST)
        if form.is_valid():
            user = form.save()
            login(request, user)
            return redirect('dashboard')
    else:
        form = UserCreationForm()
    return render(request, 'register.html', {'form': form})

def login_user(request):
    if request.method == 'POST':
        form = AuthenticationForm(request, data=request.POST)
        if form.is_valid():
            user = authenticate(request, username=form.cleaned_data['username'], password=form.cleaned_data['password'])
            if user is not None:
                login(request, user)
                return redirect('dashboard')
    else:
        form = AuthenticationForm()
    return render(request, 'login.html', {'form': form})

def logout_user(request):
    logout(request)
    return redirect('login')
    
from django.core.paginator import Paginator

@login_required
def all_emails(request):
    user_emails = EmailEntry.objects.filter(user=request.user)
    total_user_emails = user_emails.count()

    paginator = Paginator(user_emails, 6)  # Emails 6 kwa kila page
    page_number = request.GET.get('page')
    page_obj = paginator.get_page(page_number)

    return render(request, 'all_emails.html', {
        'page_obj': page_obj,
        'total_user_emails': total_user_emails
    })

@login_required
def send_selected_emails(request):
    selected_ids = request.POST.getlist('selected_emails')
    selected_emails = EmailEntry.objects.filter(id__in=selected_ids)

    for email in selected_emails:
        message = render_to_string('email_template.html', {
            'email': email,
            'message_body': "This is a test email body. Customize as needed."
        })
        plain_message = strip_tags(message)
        send_mail('Subject: Important Message', plain_message, settings.DEFAULT_FROM_EMAIL, [email.email], html_message=message)

    return redirect('dashboard')

@login_required
def send_to_all_emails(request):
    emails = EmailEntry.objects.all()

    for email in emails:
        message = render_to_string('email_template.html', {
            'email': email,
            'message_body': "This is a test email body. Customize as needed."
        })
        plain_message = strip_tags(message)
        send_mail('Subject: Important Message', plain_message, settings.DEFAULT_FROM_EMAIL, [email.email], html_message=message)

    return redirect('dashboard')






from django.contrib import messages  # Hakikisha imewekwa juu


from django.core.paginator import Paginator
from django.db.models import Q
from datetime import datetime
@login_required
def send_custom_message(request):
    if request.method == 'POST':
        subject = request.POST.get('subject')
        message_body = request.POST.get('message_body')
        send_option = request.POST.get('send_option')
        selected_ids = request.POST.get('selected_ids').split(',') if request.POST.get('selected_ids') else []

        if send_option == 'selected':
            emails = EmailEntry.objects.filter(id__in=selected_ids)
        else:
            emails = EmailEntry.objects.all()

        sent_count = 0
        for email in emails:
            message = render_to_string('email_template.html', {
                'email': email,
                'message_body': message_body,
                'subject': subject,
                'current_year': datetime.now().year
            })
            plain_message = strip_tags(message)
            send_mail(subject, plain_message, settings.DEFAULT_FROM_EMAIL, [email.email], html_message=message)
            sent_count += 1

        messages.success(request, f'Message sent to {sent_count} email{"s" if sent_count != 1 else ""}.')
        return redirect('dashboard')


    else:
        search_query = request.GET.get('search', '')
        user_emails = EmailEntry.objects.filter(user=request.user)

        if search_query:
            user_emails = user_emails.filter(
                Q(region__icontains=search_query) |
                Q(district__icontains=search_query) |
                Q(school_name__icontains=search_query) |
                Q(email__icontains=search_query) |
                Q(phone_number__icontains=search_query)
            )

        paginator = Paginator(user_emails, 3)
        page_number = request.GET.get('page')
        page_obj = paginator.get_page(page_number)

        return render(request, 'compose_message.html', {
            'user_emails': page_obj,
            'search_query': search_query,
        })


    
    
    

from django.contrib.auth.models import User  # Make sure this is imported

@login_required
def all_emails_by_user(request):
    users = User.objects.all()
    paginated_users = []

    for user in users:
        user_emails = EmailEntry.objects.filter(user=user)
        paginator = Paginator(user_emails, 4)
        page_number = request.GET.get(f'page_{user.id}')
        page_obj = paginator.get_page(page_number)

        paginated_users.append({
            'user': user,
            'page_obj': page_obj,
            'email_count': user_emails.count(),  # Count of emails
        })

    return render(request, 'all_emails_by_user.html', {'paginated_users': paginated_users})


    
from django.core.paginator import Paginator
from django.http import HttpResponse
from django.template.loader import render_to_string
from xhtml2pdf import pisa
from openpyxl import Workbook
from django.db.models import Q
import csv
import io

from .models import EmailEntry


def export_to_excel(entries):
    wb = Workbook()
    ws = wb.active
    ws.append(['School Name', 'Region', 'District', 'Email', 'Phone Number'])
    for entry in entries:
        ws.append([entry.school_name, entry.region, entry.district, entry.email, entry.phone_number])
    response = HttpResponse(
        content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
    )
    response['Content-Disposition'] = 'attachment; filename="email_entries.xlsx"'
    wb.save(response)
    return response


def render_to_pdf(template_src, context_dict={}):
    html_string = render_to_string(template_src, context_dict)
    result = io.BytesIO()
    pdf = pisa.pisaDocument(io.BytesIO(html_string.encode("UTF-8")), result)
    if not pdf.err:
        return HttpResponse(result.getvalue(), content_type='application/pdf')
    return None


def emailentry_list(request):
    email        = request.GET.get('email', '')
    phone        = request.GET.get('phone', '')
    region       = request.GET.get('region', '')
    district     = request.GET.get('district', '')
    school_name  = request.GET.get('school_name', '')
    order_by_col = request.GET.get('order_by', 'school_name')
    order_dir    = request.GET.get('order_dir', 'asc')
    contact_type = request.GET.get('contact_type', '')
    export       = request.GET.get('export', '')

    # Build base queryset
    entries = EmailEntry.objects.all()

    # Standard filters
    if email:
        entries = entries.filter(email__icontains=email)
    if phone:
        entries = entries.filter(phone_number__icontains=phone)
    if region:
        entries = entries.filter(region=region)
    if district:
        entries = entries.filter(district=district)
    if school_name:
        entries = entries.filter(school_name__icontains=school_name)

    # Contact-type filters
    if contact_type == 'email_only':
        entries = entries.filter(email__isnull=False).exclude(email='') \
                         .filter(Q(phone_number__isnull=True) | Q(phone_number=''))
    elif contact_type == 'phone_only':
        entries = entries.filter(phone_number__isnull=False).exclude(phone_number='') \
                         .filter(Q(email__isnull=True) | Q(email=''))
    elif contact_type == 'both':
        entries = entries.filter(email__isnull=False).exclude(email='') \
                         .filter(phone_number__isnull=False).exclude(phone_number='')
    elif contact_type == 'one_only':
        entries = entries.filter(
            (Q(email__isnull=False) & ~Q(email='') & (Q(phone_number__isnull=True) | Q(phone_number=''))) |
            (Q(phone_number__isnull=False) & ~Q(phone_number='') & (Q(email__isnull=True) | Q(email='')))
        )

    # Apply sorting
    prefix = '-' if order_dir == 'desc' else ''
    entries = entries.order_by(f'{prefix}{order_by_col}')

    # Export handlers
    if export == 'csv':
        resp = HttpResponse(content_type='text/csv')
        resp['Content-Disposition'] = 'attachment; filename="email_entries.csv"'
        writer = csv.writer(resp)
        writer.writerow(['School Name', 'Region', 'District', 'Email', 'Phone Number'])
        for e in entries:
            writer.writerow([e.school_name, e.region, e.district, e.email, e.phone_number])
        return resp

    if export == 'pdf':
        return render_to_pdf('pdf_template.html', {'entries': entries})

    if export == 'excel':
        return export_to_excel(entries)

    # Pagination
    paginator = Paginator(entries, 10)
    page_obj = paginator.get_page(request.GET.get('page'))

    # Dropdown options
    regions   = EmailEntry.objects.values_list('region', flat=True).distinct()
    districts = EmailEntry.objects.values_list('district', flat=True).distinct()

    # Pass everything to template
    context = {
        'entries': page_obj,
        'page_obj': page_obj,
        'filters': {
            'email': email,
            'phone': phone,
            'region': region,
            'district': district,
            'school_name': school_name,
            'contact_type': contact_type,
        },
        'order_by': order_by_col,
        'order_dir': order_dir,
        'regions': regions,
        'districts': districts,
    }
    return render(request, 'emailentry_list.html', context)





@login_required
def add_sms_message(request):
    if request.method == 'POST':
        form = SMSMessageForm(request.POST)
        if form.is_valid():
            sms = form.save(commit=False)
            sms.user = request.user
            sms.save()

             # Chagua random success message na badilisha {username}
            # message_text = random.choice(success_messages).format(username=request.user.username)
            # messages.success(request, message_text)
            msg = random.choice(success_messages).format(username=request.user.username)
            messages.success(request, msg)


            # Badala ya redirect, render template moja kwa moja
            form = SMSMessageForm()  # reset form baada ya save
            return render(request, 'add_sms_message.html', {'form': form})
    else:
        form = SMSMessageForm()
    return render(request, 'add_sms_message.html', {'form': form})



@login_required
def sms_message_list(request):
    search_query = request.GET.get('search', '')
    category_filter = request.GET.get('category', '')
    messages_qs = SMSMessage.objects.all()

    # Count zote kabla ya filter
    total_messages = messages_qs.count()
    spam_count = SMSMessage.objects.filter(category='spam').count()
    ham_count = SMSMessage.objects.filter(category='ham').count()

    # Apply filters
    if search_query:
        messages_qs = messages_qs.filter(message__icontains=search_query)
    if category_filter in ['spam', 'ham']:
        messages_qs = messages_qs.filter(category=category_filter)

    # Count baada ya filter
    filtered_count = messages_qs.count()

    # Pagination
    paginator = Paginator(messages_qs.order_by('-date_added'), 10)
    page_number = request.GET.get('page')
    page_obj = paginator.get_page(page_number)

    # Export
    export_format = request.GET.get('export', '')
    if export_format == 'csv':
        return export_sms_csv(messages_qs)
    elif export_format == 'excel':
        return export_sms_excel(messages_qs)
    elif export_format == 'pdf':
        return export_sms_pdf(messages_qs)

    return render(request, 'sms_message_list.html', {
        'page_obj': page_obj,
        'search_query': search_query,
        'category_filter': category_filter,
        'total_messages': total_messages,
        'filtered_count': filtered_count,
        'spam_count': spam_count,
        'ham_count': ham_count,
    })


def export_sms_csv(queryset):
    response = HttpResponse(content_type='text/csv')
    response['Content-Disposition'] = 'attachment; filename="sms_messages.csv"'
    writer = csv.writer(response)
    writer.writerow(['Category', 'Message'])
    for msg in queryset:
        writer.writerow([msg.category, msg.message])
    return response

def export_sms_excel(queryset):
    wb = Workbook()
    ws = wb.active
    ws.append(['Category', 'Message'])
    for msg in queryset:
        ws.append([msg.category, msg.message])
    response = HttpResponse(
        content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
    )
    response['Content-Disposition'] = 'attachment; filename="sms_messages.xlsx"'
    wb.save(response)
    return response

def export_sms_pdf(queryset):
    html = render_to_string('sms_message_pdf.html', {'messages': queryset})
    result = io.BytesIO()
    pdf = pisa.pisaDocument(io.BytesIO(html.encode("UTF-8")), result)
    if not pdf.err:
        return HttpResponse(result.getvalue(), content_type='application/pdf')
    return HttpResponse("Error generating PDF", status=500)
